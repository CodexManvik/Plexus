from __future__ import annotations

import io
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List


def _extract_docx_text(file_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as archive:
        xml_content = archive.read("word/document.xml")
    root = ET.fromstring(xml_content)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs: List[str] = []
    for para in root.findall(".//w:p", ns):
        text = "".join(node.text for node in para.findall(".//w:t", ns) if node.text)
        if text.strip():
            paragraphs.append(text.strip())
    return "\n".join(paragraphs)


def _extract_pdf_text(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return ""

    reader = PdfReader(io.BytesIO(file_bytes))
    texts: List[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            texts.append(page_text.strip())
    return "\n".join(texts)


def _extract_xlsx_text(file_bytes: bytes) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception:
        return ""

    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    fragments: List[str] = []
    for sheet in workbook.worksheets:
        fragments.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell not in (None, "")]
            if cells:
                fragments.append(" | ".join(cells))
    return "\n".join(fragments)


def extract_document_text(file_name: str, content_type: str | None, file_bytes: bytes) -> str:
    """
    Parses document bytes and returns extracted text for supported formats.
    Supported: TXT, PDF, DOCX, XLSX.
    """
    suffix = Path(file_name).suffix.lower()
    ctype = (content_type or "").lower()

    if suffix in {".txt", ".csv", ".json"} or ctype.startswith("text/"):
        return file_bytes.decode("utf-8", errors="ignore").strip()

    if suffix == ".docx" or "wordprocessingml" in ctype:
        return _extract_docx_text(file_bytes).strip()

    if suffix == ".pdf" or ctype == "application/pdf":
        return _extract_pdf_text(file_bytes).strip()

    if suffix in {".xlsx", ".xlsm"} or "spreadsheet" in ctype:
        return _extract_xlsx_text(file_bytes).strip()

    return file_bytes.decode("utf-8", errors="ignore").strip()
